from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = r'D:\CN Project\Project Title and Abstract Revised.docx'
d = Document()
s = d.sections[0]
s.page_width, s.page_height = Inches(8.27), Inches(11.69)
s.top_margin = s.bottom_margin = Inches(.72)
s.left_margin = s.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    st = d.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0,0,0)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(7)
d.styles['Normal'].paragraph_format.line_spacing = 1.08
d.styles['Title'].font.size = Pt(22)
d.styles['Heading 1'].font.size = Pt(15)
d.styles['Heading 2'].font.size = Pt(12)
def p(t): d.add_paragraph(t)
def h(t): d.add_heading(t,1)
def page(t): d.add_page_break(); h(t)

p('COMPUTER NETWORKS LAB  |  BCE308P')
d.add_heading('Quantifying Structural Exposure in Graph Based Network Intrusion Detection',0)
p('Chelintha Lingam T  |  24BCE2098\nAre Sanjana Reddy  |  24BDS0187')
h('Abstract')
p('Graph neural networks (GNNs) can exploit relationships among network flows, but their reported advantage over flow-level classifiers may depend on the graph context available during training and evaluation. Prior research has examined leakage-aware splitting, deployment-realistic temporal evaluation and topology interventions. This project focuses on the relationship between explicitly measured train/test structural exposure and the performance gap between GraphSAGE and a comparable multilayer perceptron (MLP).')
p('Using CICIDS2017 labelled flows with host and time metadata, we will compare Random Forest, an approximately parameter-matched MLP and GraphSAGE on identical prediction records and flow features. Each graph node will represent a flow; edges will connect flows sharing a communicating host within a predefined time window. We will evaluate random transductive, separate-graph inductive and chronological protocols with explicit rules governing feature and topology access. Preprocessing and model selection will use training and validation data only.')
p('The study will quantify host reuse, cross-split connectivity and one- and two-hop exposure to training flows. A fixed-record-split experiment will vary cross-split connectivity and measure the resulting GraphSAGE–MLP macro-F1 gap. Degree-controlled connectivity comparisons, degree-preserving topology rewiring and a GraphSAGE control with neighbour messages disabled will help distinguish sensitivity to exposure from changes in available context or network structure. We will report class coverage, false-positive rate, attack recall, PR-AUC and uncertainty across repeated runs.')
p('The project does not assume that overlap constitutes leakage or that GraphSAGE gains are inflated. It tests whether and under which information-access conditions structural exposure influences the observed advantage. Its intended contribution is a reproducible exposure-aware evaluation and controlled empirical evidence, with conclusions limited to the tested data, graph construction and protocols.')
p('Keywords: network intrusion detection; GraphSAGE; structural exposure; evaluation methodology; CICIDS2017.')

page('Research question and contribution')
p('Research question: How does measured train/test structural exposure affect GraphSAGE performance relative to a flow-only MLP when prediction records, features and evaluation conditions are controlled?')
p('The proposed contribution is the joint analysis of explicit exposure metrics, fixed-split connectivity interventions and model controls. Leakage-aware evaluation, MLP comparisons and topology rewiring are established components [1–5]; they are supporting methods rather than standalone novelty claims.')
h('Graph and prediction task')
p('Begin with binary benign-versus-attack flow classification. Use a unique flow record as each graph node and retain the same numeric flow-feature vector for all models. Host identifiers and timestamps are graph/split metadata, not initial classifier inputs. Use a fixed time window and bounded neighbour rule without label-based edge construction. Record whether an edge represents a shared source, destination or either endpoint. Do not form unrestricted host-sharing cliques.')
h('Evaluation protocols')
p('Random transductive: randomly partition target records, allow a declared shared graph with unlabelled held-out features visible during message passing, and restrict supervised loss to training labels. This is a diagnostic transductive setting, not automatically evidence of misconduct or leakage.')
p('Separate-graph inductive: partition records before building each graph and prohibit training-time access to held-out records and their features. GraphSAGE parameters transfer to a separate test graph. Hosts may recur; host-disjoint evaluation is a distinct optional protocol, not an implied property of induction.')
p('Chronological: train on earlier traffic and evaluate on later traffic. Fit preprocessing on training only. Use features available at flow completion and permit only graph context available at the declared prediction time. No message-passing path may bring future-flow features into an earlier prediction. Apply a boundary gap if windows or sessions would otherwise span partitions. Report attack families absent from training.')
h('Exposure metrics')
p('Host reuse = number of distinct test hosts also present in training divided by the number of distinct test hosts. Cross-split connectivity = number of training–test edges divided by the number of eligible graph edges; state the directed/undirected convention. K-hop exposure = fraction of test flow nodes with at least one training flow reachable through eligible message-passing paths within K hops, for K = 1 and 2. Audit validation exposure separately and disclose the graph on which each metric is computed.')

page('Controlled experiments and reporting')
h('Fixed split connectivity experiment')
p('Freeze train, validation and test flow IDs, features, labels and model-selection settings. Vary the retained proportion of training–test connectivity across predefined levels, such as 0%, 25%, 50%, 75% and 100%, and report measured exposure rather than assuming those levels equal exposure percentages. Keep validation isolated from the test sweep. Retrain GraphSAGE at each condition with matched seed schedules; use the same frozen test records for the MLP comparison.')
p('Simple edge deletion changes degree, isolation and graph context as well as exposure. Include constrained edge swaps or another explicit degree-controlled construction where feasible. Preserve time eligibility and split restrictions. Log any unsatisfied constraints and do not interpret deletion-only results as isolating the causal effect of overlap. The random/transductive sweep is an explanatory diagnostic; deployment claims must rely on the chronological protocol.')
h('Topology and architecture controls')
p('Compare the original topology against repeated degree-preserving rewirings that obey the same time eligibility, split boundaries and exposure constraints. Quantify the achieved change, degree preservation and exposure preservation. Avoid label-conditioned rewiring. Retrain on each topology. Disabling neighbour messages in GraphSAGE provides an additional architectural control; retain its learned self transformations and disclose inactive parameters.')
p('Approximately match the MLP and GraphSAGE parameter counts, report active counts, and provide comparable hyperparameter-search budgets and stopping rules. Parameter matching alone does not imply equivalent expressive capacity. Keep numerical preprocessing, loss weighting and target labels consistent across neural baselines.')
h('Results needed for a paper')
p('Report macro-F1, attack recall, benign false-positive rate and binary PR-AUC, with per-class support for multiclass extensions. Use at least five seeds for the principal comparisons and report paired differences. Bootstrap uncertainty at a suitable host/session/time-block level when possible; random-seed variability alone is not a sampling-confidence interval.')
p('Produce three principal outputs: protocol comparison; measured exposure versus GraphSAGE–MLP performance gap; and original versus controlled-rewired topology at matched exposure. Include split manifests, duplicate audits, graph statistics, parameter counts, runtime/memory measurements and reproducible configuration files.')
p('For a stronger general claim, replicate on a second dataset with host/time metadata and test graph-window and neighbour-limit sensitivity. A single-dataset study must explicitly scope its conclusions. Negative or null results are informative if the experiment has adequate controls and uncertainty reporting.')

page('Related work and research positioning')
p('Friji, Olivereau and Sarkiss [1] propose a flow graph, identify evaluation leakage, compare random and IP-based validation and include MLP results. Their study establishes that graph representation and split design are already active research topics. The present project focuses on measured exposure and a fixed-split intervention curve rather than proposing a new graph architecture.')
p('Van Langendonck, Castell-Uroz and Barlet-Ros [2] address practical spatio-temporal graph intrusion detection. The published PPT-GNN abstract explicitly identifies overlapping-subgraph evaluation as a source of performance overestimation and leakage. This work therefore cannot be classified as having no leakage analysis. Its published 2025 version should be distinguished from the 2024 preprint.')
p('Dey and Quaium [3] introduce leakage-safe multi-view graph construction on Edge-IIoT data. Their author-uploaded manuscript discusses training-only graph construction, held-out classification, multiple graph views, robustness tests and view overlap. View overlap is different from train/test host reuse or message-passing exposure. The project builds on this distinction without claiming that leakage-safe construction is new.')
p('Chan, Kim, Jo and Arifuzzaman [4] examine structural sensitivity on ToN-IoT and CICIDS2017 using several GNNs including GraphSAGE. They test triangle removal, bridge removal and degree-preserving rewiring under stratified edge-level cross-validation. Their study is essential prior work for topology controls. The proposed distinction is explicitly measuring and manipulating train/test exposure while examining the gap to a flow-only MLP.')
p('Moczkodan and Ragab [5] evaluate temporal sequence models on CICIDS2017 using random and leakage-free splits and padding ablations. This is methodological evidence about evaluation artifacts, not a graph-inductive experiment; a literature matrix must not treat those settings as equivalent.')
p('The dataset publication and official release [6] document distinct attack schedules across the capture days. Temporal splits can change attack coverage and class composition. These shifts must be audited before attributing a temporal performance drop to structural exposure.')
p('No direct match to the complete proposed experiment was identified in the targeted sources reviewed. This is a bounded literature assessment, not evidence that no such work exists. The project should avoid claims of being the first overlap-aware IDS benchmark or the first structural intervention study.')

page('Publication and patent assessment')
h('Conference paper potential')
p('Assessment: plausible empirical conference or workshop contribution, conditional on implementation and informative results. The proposed framework is not yet a completed paper. Its value must come from a controlled, reproducible finding about when graph context helps and how evaluation settings change the apparent advantage.')
p('Conference criteria emphasize novelty, practical significance, execution and presentation [7]. Security venues also explicitly accept work that verifies, refutes or refines earlier findings [8]. These criteria support the suitability of rigorous evaluation research, but do not establish that this particular proposal would be accepted. The cited 2026 calls illustrate criteria, not currently open submission opportunities.')
p('A specialist workshop or applied networking/security conference is a realistic initial direction if the principal controls are implemented. A stronger full-paper submission would benefit from a second dataset, robust uncertainty estimates, sensitivity analyses and a reproducible artifact. No venue tier or acceptance probability is asserted.')
h('Patent assessment for India')
p('Assessment: patentability is not established, and the present benchmarking proposal is a weak patent candidate. India requires novelty, inventive step and industrial applicability, with no applicable statutory exclusion [9]. Section 3(k) excludes mathematical methods, algorithms and computer programs per se [10]. The final CRI guidelines provide the relevant examination framework for computer-related inventions [11]. Software is not categorically unpatentable; the substance of the claimed technical contribution matters.')
p('A benchmarking workflow and performance-gap measurements do not by themselves demonstrate a non-obvious technical invention. A separate operational mechanism that changes live IDS processing and produces a substantiated technical improvement would need its own invention definition and claim-by-claim prior-art review. Simply adding a deployment interface or renaming the framework would not establish inventive step.')
p('A related patent publication [12] covers NIDS training and inference using overlapping flow windows, image conversion and feedback-based overlap adjustment. It concerns temporal window overlap, not the proposed graph-exposure experiment. It is background prior art rather than a determination that the proposal infringes or lacks novelty. This review is not a comprehensive patent search, freedom-to-operate opinion or legal patentability opinion.')
p('Recommended route: implement the empirical study and pursue a paper. If a distinct operational invention emerges, consult the institution’s IP cell or a registered patent professional before public disclosure; disclosure can affect novelty. Assess any such invention separately from the present evaluation framework.')

page('References')
refs = [
('[1] H. Friji, A. Olivereau and M. Sarkiss. Efficient Network Representation for GNN-based Intrusion Detection. ACNS 2023; arXiv:2310.05956.', 'https://arxiv.org/html/2310.05956v1'),
('[2] L. Van Langendonck, I. Castell-Uroz and P. Barlet-Ros. PPT-GNN: A Practical Pretrained Spatio-Temporal Graph Neural Network for Network Security. IEEE EuroS&P Workshops, 2025. DOI: 10.1109/EuroSPW67616.2025.00026.', 'https://ieeexplore.ieee.org/document/11129530/'),
('[3] S. Dey and A. Quaium. Leakage-Safe Multi-View Graph Construction for Robust Intrusion Detection in Edge-IIoT Networks. NSysS 2025. DOI: 10.1145/3777555.3777572. Author-uploaded manuscript.', 'https://www.researchgate.net/publication/398782840_Leakage-Safe_Multi-View_Graph_Construction_for_Robust_Intrusion_Detection_in_Edge-IIoT_Networks'),
('[4] K. O. Chan, Y. Kim, J.-Y. Jo and S. Arifuzzaman. Structural Sensitivity of Graph Neural Networks for Intrusion Detection: An Intervention Study. IEEE Access, 2026. DOI: 10.1109/ACCESS.2026.3703442. Author manuscript reviewed.', 'https://doi.org/10.1109/ACCESS.2026.3703442'),
('[5] Z. Moczkodan and H. Ragab. Do Transformers Actually Help Intrusion Detection? A Temporal Sequence Evaluation on CIC-IDS2017. arXiv:2606.11098v2, 2026.', 'https://arxiv.org/html/2606.11098v2'),
('[6] Canadian Institute for Cybersecurity, UNB. Intrusion detection evaluation dataset CIC-IDS2017. Dataset paper: I. Sharafaldin, A. H. Lashkari and A. A. Ghorbani, ICISSP 2018.', 'https://www.unb.ca/cic/datasets/ids-2017.html'),
('[7] ACM CCS 2026. Call for Papers. Research and artifact review criteria.', 'https://www.sigsac.org/ccs/CCS2026/call-for/call-for-papers.html'),
('[8] USENIX Security 2026. Call for Papers. Replication and open-science provisions.', 'https://www.usenix.org/conference/usenixsecurity26/call-for-papers'),
('[9] IP India. Basics of Patents. Patentability and inventive-step requirements.', 'https://ipindia.gov.in/basics-of-patents'),
('[10] IP India. Patents Act 1970, Section 3. Statutory exclusions.', 'https://www.ipindia.gov.in/acts/patent-act-1970/section-3'),
('[11] IP India. Guidelines for Examination of Computer Related Inventions 2025. Final guidelines linked from the official resources index.', 'https://ipindia.gov.in/resource/patents-resources-guidelines'),
('[12] WO2024180374A1. Method for training and using a network intrusion detection system. Published 6 September 2024.', 'https://patents.google.com/patent/WO2024180374A1/en'),
]
for title,url in refs:
    para=d.add_paragraph(title)
    para.paragraph_format.space_after=Pt(2)
    para=d.add_paragraph()
    link=OxmlElement('w:hyperlink')
    link.set(qn('r:id'), para.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run=OxmlElement('w:r')
    props=OxmlElement('w:rPr')
    size=OxmlElement('w:sz'); size.set(qn('w:val'),'18'); props.append(size)
    under=OxmlElement('w:u'); under.set(qn('w:val'),'single'); props.append(under)
    run.append(props)
    text=OxmlElement('w:t'); text.text='Open source'; run.append(text)
    link.append(run); para._p.append(link)
    para.paragraph_format.space_after=Pt(8)
    for r in para.runs:r.font.size=Pt(9)
p('Sources checked on 8 October 2026. Proposal language describes planned experiments; no experimental findings are claimed.')
d.save(OUT)
print(OUT)
